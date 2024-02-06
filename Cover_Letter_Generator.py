from docx import Document
from docx2pdf import convert
from docx.shared import Pt, RGBColor
import os
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

def replace_text_in_docx(input_path, output_path, replacements):
    doc = Document(input_path)
    
    # Replace in paragraphs
    for old_text, new_text in replacements.items():
        for paragraph in doc.paragraphs:
            if old_text in paragraph.text:
                for run in paragraph.runs:
                    if old_text in run.text:
                        run.text = run.text.replace(old_text, new_text)
                        run.font.size = Pt(10)  # Set the desired font size (adjust as needed)
                        run.font.name = 'Avenir Next Regular'  # Set the desired font name (adjust as needed)
                        run.font.color.rgb = RGBColor(0, 0, 0) # RGB values for black

    # Replace in text boxes
    for shape in doc.inline_shapes:
        if shape.type == WD_INLINE_SHAPE.TEXT_BOX:
            text = shape.text_frame.text
            print("Text in text box:", text)
            for old_text, new_text in replacements.items():
                if old_text in text:
                    text = text.replace(old_text, new_text)
            shape.text_frame.text = text
    
    doc.save(output_path)
    convert(output_path)
    
    filename = os.path.basename(input_path)
    type = filename.split('.')[0].split(' - ')[0]
    print(f"{type} for {Company} has been generated.")

Job_Title = input("Enter job title: ")
Company = input("Enter company name: ")
Companys = f"{Company}'s"

# Input and output file paths
input_file_path = "C:/example/directory/Cover Letter.docx" 
output_file_path = f"C:/example/directory/Cover Letter - {Company}.docx"
input_file_path2 = "C:/example/directory/CV.docx" 
output_file_path2 = f"C:/example/directory/CV - {Company}.docx"

# Define the replacements as a dictionary
replacements = {
    'xxx': Job_Title,
    'yyy': 'job',
    'zzz': Company,
    "zzz's": Companys 
}

# Replace the specified text in the Word document
replace_text_in_docx(input_file_path, output_file_path, replacements)
replace_text_in_docx(input_file_path2, output_file_path2, replacements)

